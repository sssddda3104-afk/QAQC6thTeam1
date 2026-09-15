"""Convert the report displayed in Streamlit to an in-memory Word document."""
from io import BytesIO
import re


def build_report_docx(text, title, metadata):
    # Lazy imports keep the dashboard usable before export dependencies are installed.
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from markdown_it import MarkdownIt

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number']:
        style = doc.styles[name]
        style.font.name = 'Malgun Gothic'
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Malgun Gothic')
        style.font.color.rgb = RGBColor.from_string('172B4D')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    doc.styles['Title'].font.size = Pt(21)
    doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in [('Heading 1', 15), ('Heading 2', 13), ('Heading 3', 12)]:
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    doc.add_paragraph(title, 'Title')
    for key, value in metadata.items():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f'{key}  ').bold = True
        p.add_run(str(value))
    doc.add_paragraph()

    def inline(paragraph, children):
        bold = italic = False
        for token in children or []:
            if token.type == 'strong_open': bold = True
            elif token.type == 'strong_close': bold = False
            elif token.type == 'em_open': italic = True
            elif token.type == 'em_close': italic = False
            elif token.type in ('softbreak', 'hardbreak'):
                paragraph.add_run().add_break()
            elif token.type in ('text', 'code_inline', 'html_inline'):
                run = paragraph.add_run(token.content)
                run.bold, run.italic = bold, italic
            elif token.type == 'image':
                paragraph.add_run(token.content)  # Never fetch URLs from model output.

    tokens = MarkdownIt('commonmark', {'html': False}).enable('table').parse(text)
    paragraph = None
    table = None
    cells = []
    cell_index = -1
    in_header = False
    lists = []
    item_pending = False
    for i, token in enumerate(tokens):
        kind = token.type
        if kind == 'table_open':
            count = 0
            for future in tokens[i+1:]:
                if future.type == 'tr_close': break
                if future.type == 'th_open': count += 1
            table = doc.add_table(rows=0, cols=max(1, count))
            table.style = 'Table Grid'
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = ([1.0, 1.2, 1.2, 3.6] if count == 4 else [7.0 / max(count, 1)] * count)
            for column, width in zip(table.columns, widths): column.width = Inches(width)
        elif kind == 'thead_open': in_header = True
        elif kind == 'thead_close': in_header = False
        elif kind == 'tr_open':
            row = table.add_row()
            cells, cell_index = row.cells, -1
            # Keep each individual row intact; allow long tables across pages.
            props = row._tr.get_or_add_trPr()
            props.append(OxmlElement('w:cantSplit'))
            if in_header: props.append(OxmlElement('w:tblHeader'))
            for cell, column in zip(cells, table.columns):
                cell.width = column.width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                margins = OxmlElement('w:tcMar')
                for side in ['top', 'bottom', 'left', 'right']:
                    node = OxmlElement(f'w:{side}')
                    node.set(qn('w:w'), '100')
                    node.set(qn('w:type'), 'dxa')
                    margins.append(node)
                cell._tc.get_or_add_tcPr().append(margins)
                if in_header:
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'EAF0F8')
                    cell._tc.get_or_add_tcPr().append(shade)
        elif kind in ('th_open', 'td_open'):
            cell_index += 1
            paragraph = cells[cell_index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.15
        elif kind in ('th_close', 'td_close'):
            if in_header:
                for run in paragraph.runs: run.bold = True
            elif re.fullmatch(r'[\s\d.,%+−\-/]+', paragraph.text):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif kind == 'table_close':
            table = None
            doc.add_paragraph().paragraph_format.space_after = Pt(3)
        elif kind in ('bullet_list_open', 'ordered_list_open'):
            lists.append([kind, int(token.attrGet('start') or 1)])
        elif kind in ('bullet_list_close', 'ordered_list_close'):
            lists.pop()
        elif kind == 'list_item_open': item_pending = True
        elif kind == 'heading_open':
            # Prompts use H3; normalize report sections to Word Heading 1.
            paragraph = doc.add_paragraph(style=f'Heading {max(1, min(3, int(token.tag[1:]) - 2))}')
        elif kind == 'paragraph_open' and table is None:
            paragraph = doc.add_paragraph()
            if lists and item_pending:
                paragraph.paragraph_format.left_indent = Inches(.18 * len(lists))
                paragraph.paragraph_format.first_line_indent = Inches(-.18)
                marker = '• ' if lists[-1][0] == 'bullet_list_open' else f'{lists[-1][1]}. '
                paragraph.add_run(marker)
                lists[-1][1] += 1
                item_pending = False
        elif kind == 'inline': inline(paragraph, token.children)
        elif kind in ('fence', 'code_block'):
            doc.add_paragraph(token.content)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('AI 자동 리포트  |  ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    for run in footer.runs: run.font.size = Pt(9)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
